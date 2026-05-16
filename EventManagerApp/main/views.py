import csv
import io
import re
from datetime import datetime, time as datetime_time
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.conf import settings
from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.db import connection, transaction
from django.db.models import Count, Q
from django.http import FileResponse, Http404, HttpResponseForbidden
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme
from django.views.decorators.http import require_POST
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
import yake

from .decorators import role_required
from .forms import (
    ContractorForm,
    EmployeeCreationForm,
    EventEmployeeBindingForm,
    EmployeeUpdateForm,
    EventForm,
    FeedbackCsvUploadForm,
    LoginUserForm,
    MessageForm,
    OrderForm,
    ParticipantCsvUploadForm,
    ParticipantForm,
    PlaceForm,
    TaskForm,
)
from .models import Contractor, Employee, Event, Feedback, Message, Order, Participant, Place, Report, Task


EVENT_DURATION_GROUPS = (
    ("short", "Короткие до 3 часов"),
    ("day", "Средние от 3 до 24 часов"),
    ("multiday", "Длительные от 1 до 2 дней"),
    ("long", "Длинные 2+ дней"),
    ("unknown", "Без указанной длительности"),
)
EVENT_DURATION_GROUP_LABELS = dict(EVENT_DURATION_GROUPS)


def _fetch_event_employee_map(event_ids):
    if not event_ids:
        return {}

    placeholders = ", ".join(["%s"] * len(event_ids))
    sql = f"""
        SELECT ee.event_id, e.fullname
        FROM employee_on_event ee
        JOIN employee e ON e.e_id = ee.e_id
        WHERE ee.event_id IN ({placeholders})
        ORDER BY e.fullname
    """
    result = {event_id: [] for event_id in event_ids}
    with connection.cursor() as cursor:
        cursor.execute(sql, event_ids)
        for event_id, fullname in cursor.fetchall():
            result.setdefault(event_id, []).append(fullname)
    return result


def _fetch_selected_employee_ids(event_id):
    with connection.cursor() as cursor:
        cursor.execute("SELECT e_id FROM employee_on_event WHERE event_id = %s", [event_id])
        return [row[0] for row in cursor.fetchall()]


def _sync_event_employees(event_id, employee_ids):
    with connection.cursor() as cursor:
        cursor.execute("DELETE FROM employee_on_event WHERE event_id = %s", [event_id])
        for employee_id in employee_ids:
            cursor.execute(
                "INSERT INTO employee_on_event (e_id, event_id) VALUES (%s, %s)",
                [employee_id, event_id],
            )


def _get_event_orders(event):
    return list(Order.objects.filter(event=event).select_related("c"))


def _get_event_tasks(event):
    return list(Task.objects.filter(event=event).select_related("e", "operator"))


def _visible_tasks_queryset(user):
    queryset = Task.objects.select_related("event", "e", "operator").filter(event__in=_visible_events_queryset(user))
    if user.is_superuser or user.position in {Employee.ROLE_ADMIN, Employee.ROLE_TEAMLEAD}:
        return queryset
    if user.position == Employee.ROLE_MANAGER:
        return queryset.filter(Q(e=user) | Q(operator=user)).distinct()
    return queryset.filter(e=user)


def _visible_events_queryset(user):
    queryset = Event.objects.select_related("p", "created_by")
    if user.is_superuser or user.position in {Employee.ROLE_ADMIN, Employee.ROLE_TEAMLEAD}:
        return queryset

    access_filter = Q(employeeonevent__e=user)
    if user.position == Employee.ROLE_MANAGER:
        access_filter |= Q(created_by=user)
    return queryset.filter(access_filter).distinct()


def _decode_csv_file(uploaded_file):
    data = uploaded_file.read()
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8-sig", errors="replace")


def _normalize_csv_header(value):
    return (value or "").strip().lower().replace(" ", "_").replace("-", "_")


def _pick_csv_value(row, *names):
    normalized = {_normalize_csv_header(key): value for key, value in row.items()}
    for name in names:
        value = normalized.get(_normalize_csv_header(name))
        if value is not None:
            return (value or "").strip()
    return ""


def _pick_csv_value_by_keywords(row, *keywords):
    normalized_keywords = [_normalize_csv_header(keyword) for keyword in keywords]
    for key, value in row.items():
        normalized_key = _normalize_csv_header(key)
        if all(keyword in normalized_key for keyword in normalized_keywords):
            return (value or "").strip()
    return ""


def _parse_participant_gender(value):
    normalized = (value or "").strip().lower()
    if normalized in {"male", "m", "м", "муж", "мужской"}:
        return Participant.GENDER_MALE
    if normalized in {"female", "f", "ж", "жен", "женский"}:
        return Participant.GENDER_FEMALE
    return None


def _looks_like_email(value):
    return "@" in (value or "")


def _normalize_phone(value):
    raw = (value or "").strip().replace("\ufeff", "")
    if not raw:
        return ""

    if raw.startswith('="') and raw.endswith('"'):
        raw = raw[2:-1].strip()
    if raw.startswith("'"):
        raw = raw[1:].strip()

    numeric_candidate = raw.replace(" ", "").replace(",", ".")
    if re.fullmatch(r"[+]?\d+(?:\.\d+)?(?:e[+-]?\d+)?", numeric_candidate, flags=re.IGNORECASE):
        try:
            decimal_value = Decimal(numeric_candidate.lstrip("+"))
            if decimal_value == decimal_value.to_integral_value():
                raw = str(decimal_value.quantize(Decimal(1)))
        except (InvalidOperation, ValueError):
            pass

    raw = re.sub(r"\.0+$", "", raw.strip())
    digits = re.sub(r"\D", "", raw)
    if not digits:
        return raw

    if len(digits) == 10:
        return f"+7{digits}"
    if len(digits) == 11 and digits.startswith("8"):
        return f"+7{digits[1:]}"
    if len(digits) == 11 and digits.startswith("7"):
        return f"+{digits}"
    if raw.strip().startswith("+"):
        return f"+{digits}"
    return digits


def _looks_like_phone(value):
    normalized = _normalize_phone(value)
    return len(re.sub(r"\D", "", normalized)) >= 5 and not _looks_like_email(value)


def _make_participant(event, fullname, gender="", phone="", email=""):
    return Participant(
        event=event,
        fullname=(fullname or "").strip(),
        gender=_parse_participant_gender(gender),
        phone=_normalize_phone(phone),
        email=(email or "").strip(),
        attended=False,
    )


def _current_local_datetime_text():
    return timezone.localtime(timezone.now()).strftime("%d.%m.%Y %H:%M")


def _parse_task_deadline(value):
    raw = (value or "").strip()
    if not raw:
        return None

    datetime_formats = (
        "%d.%m.%Y %H:%M",
        "%d.%m.%Y %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d %H:%M:%S",
    )
    date_formats = ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y")

    for date_format in datetime_formats:
        try:
            parsed = datetime.strptime(raw, date_format)
            return timezone.make_aware(parsed, timezone.get_current_timezone()) if timezone.is_naive(parsed) else parsed
        except ValueError:
            continue

    for date_format in date_formats:
        try:
            parsed_date = datetime.strptime(raw, date_format).date()
            parsed = datetime.combine(parsed_date, datetime_time.max.replace(microsecond=0))
            return timezone.make_aware(parsed, timezone.get_current_timezone())
        except ValueError:
            continue

    return None


def _make_aware_local_datetime(value):
    if timezone.is_naive(value):
        value = timezone.make_aware(value, timezone.get_current_timezone())
    return timezone.localtime(value)


def _parse_event_datetime(value, end_of_day=False):
    raw = (value or "").strip()
    if not raw:
        return None

    datetime_formats = (
        "%d.%m.%Y %H:%M",
        "%d.%m.%Y %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%d/%m/%Y %H:%M:%S",
    )
    date_formats = ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y")

    for date_format in datetime_formats:
        try:
            parsed = datetime.strptime(raw, date_format)
            return _make_aware_local_datetime(parsed)
        except ValueError:
            continue

    for date_format in date_formats:
        try:
            parsed_date = datetime.strptime(raw, date_format).date()
            parsed_time = datetime_time.max.replace(microsecond=0) if end_of_day else datetime_time.min
            return _make_aware_local_datetime(datetime.combine(parsed_date, parsed_time))
        except ValueError:
            continue

    return None


def _event_start_datetime(event):
    return _parse_event_datetime(event.time)


def _event_end_datetime(event):
    return _parse_event_datetime(event.end_date, end_of_day=True)


def _event_duration_hours(event):
    start_at = _event_start_datetime(event)
    end_at = _event_end_datetime(event)
    if not start_at or not end_at or end_at < start_at:
        return None
    return round((end_at - start_at).total_seconds() / 3600, 2)


def _event_duration_group_key(event):
    duration_hours = _event_duration_hours(event)
    if duration_hours is None:
        return "unknown"
    if duration_hours <= 3:
        return "short"
    if duration_hours < 24:
        return "day"
    if duration_hours < 48:
        return "multiday"
    return "long"


def _event_matches_period(event, period_start=None, period_end=None):
    start_at = _event_start_datetime(event)
    if start_at is None:
        return period_start is None and period_end is None
    if period_start and start_at < period_start:
        return False
    if period_end and start_at > period_end:
        return False
    return True


def _task_closed_late(task):
    if task.status != Task.STATUS_DONE or not task.closed_at:
        return False
    deadline = _parse_task_deadline(task.deadline)
    if deadline is None:
        return False
    closed_at = task.closed_at
    if timezone.is_naive(closed_at):
        closed_at = timezone.make_aware(closed_at, timezone.get_current_timezone())
    return closed_at > deadline


def _task_timeliness_stats(tasks):
    checked_count = 0
    on_time_count = 0
    for task in tasks:
        if task.status != Task.STATUS_DONE or not task.closed_at:
            continue
        deadline = _parse_task_deadline(task.deadline)
        if deadline is None:
            continue
        checked_count += 1
        if not _task_closed_late(task):
            on_time_count += 1

    percent = round((on_time_count / checked_count) * 100, 2) if checked_count else 0
    return {
        "checked_count": checked_count,
        "on_time_count": on_time_count,
        "percent": percent,
    }


def _parse_feedback_rating(value):
    raw = (value or "").strip().replace(",", ".")
    if not raw:
        return None
    try:
        rating = int(float(raw))
    except ValueError:
        return None
    if 1 <= rating <= 10:
        return rating
    return None


def _extract_feedback_keywords(feedbacks, top=10):
    text = " ".join((feedback.review or "").strip() for feedback in feedbacks if feedback.review).strip()
    if not text:
        return []

    extractor = yake.KeywordExtractor(lan="ru", n=2, dedup_lim=0.9, top=top * 4)
    keywords = []
    seen = set()
    seen_token_sets = []
    for keyword, _score in extractor.extract_keywords(text):
        normalized = keyword.strip()
        if not normalized:
            continue

        key = _normalize_keyword(normalized)
        token_set = _keyword_token_set(key)
        if not key or key in seen or _is_duplicate_keyword(key, token_set, seen, seen_token_sets):
            continue

        seen.add(key)
        seen_token_sets.append(token_set)
        keywords.append(normalized)
        if len(keywords) >= top:
            break
    return keywords


def _normalize_keyword(value):
    normalized = (value or "").strip().casefold().replace("ё", "е")
    normalized = re.sub(r"[^\w\s-]", " ", normalized, flags=re.UNICODE)
    normalized = re.sub(r"[_\s-]+", " ", normalized).strip()
    return normalized


def _keyword_token_set(value):
    return {token for token in value.split() if len(token) > 2}


def _is_duplicate_keyword(candidate, candidate_tokens, seen_keywords, seen_token_sets):
    for existing in seen_keywords:
        if candidate in existing or existing in candidate:
            return True

    if not candidate_tokens:
        return False

    for existing_tokens in seen_token_sets:
        if not existing_tokens:
            continue
        overlap = candidate_tokens & existing_tokens
        if not overlap:
            continue
        if len(candidate_tokens) == 1 or len(existing_tokens) == 1:
            return True
        if len(overlap) / len(candidate_tokens | existing_tokens) >= 0.75:
            return True

    return False


def _import_participants_from_csv(event, uploaded_file):
    text = _decode_csv_file(uploaded_file)
    sample = text[:2048]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;")
    except csv.Error:
        dialect = csv.excel

    stream = io.StringIO(text)
    reader = csv.DictReader(stream, dialect=dialect)
    participants = []

    if not reader.fieldnames:
        return 0

    known_headers = {
        "fullname",
        "full_name",
        "fio",
        "name",
        "фио",
        "имя",
        "gender",
        "пол",
        "phone",
        "телефон",
        "номер_телефона",
        "email",
        "e_mail",
        "mail",
        "почта",
        "электронная_почта",
    }
    normalized_headers = {_normalize_csv_header(fieldname) for fieldname in reader.fieldnames}
    has_feedback_headers = normalized_headers.intersection(known_headers) or any(
        (
            "оцените_прошедшее_мероприятие" in header
            or "расскажите_что_вам_понравилось" in header
        )
        for header in normalized_headers
    )
    if not has_feedback_headers:
        stream.seek(0)
        simple_reader = csv.reader(stream, dialect=dialect)
        for columns in simple_reader:
            columns = [column.strip() for column in columns]
            if not any(columns):
                continue
            fullname = columns[0] if len(columns) > 0 else ""
            gender = ""
            phone = ""
            email = ""

            if len(columns) >= 4:
                if _looks_like_phone(columns[1]) or _looks_like_email(columns[2]):
                    phone = columns[1]
                    email = columns[2]
                    gender = columns[3]
                else:
                    gender = columns[1]
                    phone = columns[2]
                    email = columns[3]
            elif len(columns) == 3:
                if _looks_like_phone(columns[1]) or _looks_like_email(columns[2]):
                    phone = columns[1]
                    email = columns[2]
                else:
                    gender = columns[1]
                    phone = columns[2]
            elif len(columns) == 2:
                if _looks_like_phone(columns[1]):
                    phone = columns[1]
                else:
                    gender = columns[1]

            participants.append(_make_participant(event, fullname, gender=gender, phone=phone, email=email))
        if participants:
            Participant.objects.bulk_create(participants)
        return len(participants)

    for row in reader:
        fullname = _pick_csv_value(row, "fullname", "full_name", "fio", "name", "фио", "имя")
        phone = _pick_csv_value(row, "phone", "телефон", "номер телефона", "номер_телефона")
        email = _pick_csv_value(row, "email", "e_mail", "mail", "почта", "электронная_почта")
        gender = _pick_csv_value(row, "gender", "пол")

        if not any([fullname, phone, email, gender]):
            continue

        participants.append(
            _make_participant(event, fullname, gender=gender, phone=phone, email=email)
        )

    if participants:
        Participant.objects.bulk_create(participants)
    return len(participants)


def _import_feedback_from_csv(event, uploaded_file):
    text = _decode_csv_file(uploaded_file)
    sample = text[:2048]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;")
    except csv.Error:
        dialect = csv.excel

    stream = io.StringIO(text)
    reader = csv.DictReader(stream, dialect=dialect)
    feedback_items = []

    if not reader.fieldnames:
        return 0

    known_headers = {
        "rating",
        "оценка",
        "score",
        "оцените_прошедшее_мероприятие_от_1_до_10",
        "review",
        "отзыв",
        "comment",
        "комментарий",
        "text",
        "текст",
        "расскажите_что_вам_понравилось/не_понравилось_на_прошедшем_мероприятии",
    }
    normalized_headers = {_normalize_csv_header(fieldname) for fieldname in reader.fieldnames}
    if not normalized_headers.intersection(known_headers):
        stream.seek(0)
        simple_reader = csv.reader(stream, dialect=dialect)
        for columns in simple_reader:
            columns = [column.strip() for column in columns]
            if not any(columns):
                continue
            first = columns[0] if len(columns) > 0 else ""
            second = columns[1] if len(columns) > 1 else ""
            rating = _parse_feedback_rating(first)
            review = second
            if rating is None:
                rating = _parse_feedback_rating(second)
                review = first
            if rating is None or not review:
                continue
            feedback_items.append(Feedback(event=event, rating=rating, review=review))
        if feedback_items:
            Feedback.objects.bulk_create(feedback_items)
        return len(feedback_items)

    for row in reader:
        rating_value = _pick_csv_value(
            row,
            "rating",
            "оценка",
            "score",
            "оцените прошедшее мероприятие от 1 до 10",
        )
        if not rating_value:
            rating_value = _pick_csv_value_by_keywords(row, "оцените", "мероприятие")
        rating = _parse_feedback_rating(rating_value)
        review = _pick_csv_value(
            row,
            "review",
            "отзыв",
            "comment",
            "комментарий",
            "text",
            "текст",
            "расскажите что вам понравилось/не понравилось на прошедшем мероприятии",
        )
        if not review:
            review = _pick_csv_value_by_keywords(row, "расскажите", "понравилось")
        if rating is None or not review:
            continue
        feedback_items.append(Feedback(event=event, rating=rating, review=review))

    if feedback_items:
        Feedback.objects.bulk_create(feedback_items)
    return len(feedback_items)


def _split_employees_by_role(event):
    employees = Employee.objects.filter(employeeonevent__event=event).distinct()
    grouped = {
        Employee.ROLE_TEAMLEAD: [],
        Employee.ROLE_MANAGER: [],
        Employee.ROLE_ASSISTANT: [],
    }
    for employee in employees:
        grouped.setdefault(employee.position, []).append(employee.fullname or employee.login)
    return grouped


def _expense_total(orders):
    return sum(order.quantity * order.price for order in orders)


def _planned_budget(event):
    return float(event.planned_budget or 0)


def _budget_deviation_percent(planned_budget, actual_expenses):
    if planned_budget <= 0:
        return None
    return round(((actual_expenses - planned_budget) / planned_budget) * 100, 2)


def _budget_quality_deviation(planned_budget, actual_expenses):
    if planned_budget <= 0:
        return 0
    return (planned_budget - actual_expenses) / planned_budget


def _average_event_rating(event):
    ratings = list(Feedback.objects.filter(event=event).values_list("rating", flat=True))
    if not ratings:
        return None
    return round(sum(ratings) / len(ratings), 2)


def _attendance_ratio(event):
    participants_count = Participant.objects.filter(event=event).count()
    if not participants_count:
        return 0
    attended_count = Participant.objects.filter(event=event, attended=True).count()
    return attended_count / participants_count


def _complex_event_score(avg_rating, attendance_ratio, planned_budget, actual_expenses):
    rating_component = (avg_rating / 10) if avg_rating is not None else 0
    budget_component = _budget_quality_deviation(planned_budget, actual_expenses)
    return round(0.85 * rating_component + 0.15 * attendance_ratio + 0.05 * budget_component, 4)


def _event_quality_score(event):
    avg_rating = _average_event_rating(event)
    planned_budget = _planned_budget(event)
    actual_expenses = _expense_total(_get_event_orders(event))
    return _complex_event_score(avg_rating, _attendance_ratio(event), planned_budget, actual_expenses)


def _build_rating_histogram(feedbacks):
    counts = {rating: 0 for rating in range(1, 11)}
    for feedback in feedbacks:
        if feedback.rating in counts:
            counts[feedback.rating] += 1

    max_count = max(counts.values()) if counts else 0
    return [
        {
            "rating": rating,
            "count": count,
            "bar_percent": round((count / max_count) * 100) if max_count else 0,
        }
        for rating, count in counts.items()
    ]


def _build_rating_chart_data(rating_histogram):
    return {
        "labels": [f"{item['rating']}/10" for item in rating_histogram],
        "counts": [item["count"] for item in rating_histogram],
    }


def _score_employees(tasks, event, event_average_rating):
    employees = Employee.objects.filter(employeeonevent__event=event).distinct()
    rows = []
    for employee in employees:
        own_tasks = [task for task in tasks if task.e_id == employee.pk or task.operator_id == employee.pk]
        if not own_tasks or not isinstance(event_average_rating, (int, float)):
            score = "Н/Д"
        else:
            completed = sum(1 for task in own_tasks if task.status == Task.STATUS_DONE)
            incomplete_share = (len(own_tasks) - completed) / len(own_tasks)
            score = round(max(0, event_average_rating - incomplete_share), 2)
        rows.append((employee.fullname or employee.login, score))
    return rows


def _score_contractors(orders, event):
    contractor_rows = []
    grouped = {}
    for order in orders:
        grouped.setdefault(order.c_id, {"contractor": order.c, "amount": 0, "count": 0})
        grouped[order.c_id]["amount"] += order.quantity * order.price
        grouped[order.c_id]["count"] += 1

    for data in grouped.values():
        contractor = data["contractor"]
        score = max(3, min(10, 10 - (data["count"] - 1)))
        if event.status == Event.STATUS_FINISHED:
            conclusion = "Продолжить сотрудничество" if score >= 7 else "Провести переговоры"
        else:
            conclusion = "Оценка предварительная"
        contractor_rows.append((contractor.name, score, conclusion))
    return contractor_rows


def _average_numeric(values):
    numbers = [value for value in values if isinstance(value, (int, float))]
    if not numbers:
        return "Н/Д"
    return round(sum(numbers) / len(numbers), 2)


def _set_report_font(paragraph, size=12, bold=False, italic=False):
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic


def _add_report_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        hdr_cells[idx].text = str(header)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)
    return table


def _delete_report_file(rep_path):
    if not rep_path:
        return
    report_file = Path(rep_path)
    try:
        if report_file.exists() and report_file.is_file():
            report_file.unlink()
    except OSError:
        pass


def _build_event_report_doc(event):
    document = Document()
    orders = _get_event_orders(event)
    tasks = _get_event_tasks(event)
    participants_count = Participant.objects.filter(event=event).count()
    attended_count = Participant.objects.filter(event=event, attended=True).count()
    attendance_ratio = round((attended_count / participants_count) * 100, 2) if participants_count else 0
    grouped_employees = _split_employees_by_role(event)
    timeliness = _task_timeliness_stats(tasks)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"ОТЧЕТ О МЕРОПРИЯТИИ № {event.pk}\n").bold = True
    p.add_run((event.title or "").upper()).bold = True
    _set_report_font(p, size=13, bold=True)

    auto = document.add_paragraph()
    auto.add_run("СГЕНЕРИРОВАНО АВТОМАТИЧЕСКИ").italic = True
    _set_report_font(auto, size=12, italic=True)

    info = [
        ("Дата и время проведения мероприятия", event.time or "Не указано"),
        ("Дата окончания мероприятия", event.end_date or "Не указано"),
        ("Место проведения мероприятия", event.p.address if event.p else "Не указано"),
        ("Количество участников", participants_count),
        ("Доля пришедших участников", f"{attended_count} из {participants_count} ({attendance_ratio}%)"),
        (
            "Процент вовремя закрытых задач",
            f"{timeliness['percent']}% ({timeliness['on_time_count']} из {timeliness['checked_count']})",
        ),
    ]
    for label, value in info:
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{label} – ").bold = True
        paragraph.add_run(str(value))
        _set_report_font(paragraph)

    for label, role_key in [
        ("Руководитель(и)", Employee.ROLE_TEAMLEAD),
        ("Менеджер(ы)", Employee.ROLE_MANAGER),
        ("Ассистент(ы)", Employee.ROLE_ASSISTANT),
    ]:
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{label}:").bold = True
        _set_report_font(paragraph)
        names = grouped_employees.get(role_key) or ["Не назначены"]
        for name in names:
            item = document.add_paragraph(name)
            item.paragraph_format.left_indent = Pt(18)
            _set_report_font(item)

    linked = document.add_paragraph()
    linked.add_run("Связанные заказы").bold = True
    _set_report_font(linked)
    order_rows = []
    for order in orders:
        order_rows.append(
            [
                order.c.name,
                order.c.fullname,
                f"ORD-{order.order_id}",
                order.date,
                order.c.get_type_display() if hasattr(order.c, "get_type_display") else order.c.type,
            ]
        )
    if order_rows:
        _add_report_table(
            document,
            ["Юр лицо", "Контактное лицо", "Номер договора", "Дата заключения", "Тип договора"],
            order_rows,
        )
    else:
        empty = document.add_paragraph("Связанные заказы отсутствуют.")
        _set_report_font(empty)

    comment = document.add_paragraph()
    comment.add_run("Комментарий менеджера: ").bold = True
    comment.add_run(event.description or "Комментарий к мероприятию не добавлен.")
    _set_report_font(comment)

    feedback_title = document.add_paragraph()
    feedback_title.add_run("Обратная связь:").bold = True
    _set_report_font(feedback_title)
    feedback_rows = []
    feedback_items = Feedback.objects.filter(event=event)
    for feedback in feedback_items:
        feedback_rows.append([feedback.feedback_id, feedback.review, feedback.rating])
    if not feedback_rows:
        feedback_rows.append([1, "Данные обратной связи не собраны.", "Н/Д"])
    avg_feedback = _average_numeric([row[2] for row in feedback_rows])
    feedback_rows.append(["Среднее", "", avg_feedback])
    _add_report_table(document, ["ID", "Отзыв", "Оценка (от 1 до 10)"], feedback_rows)
    employee_scores = _score_employees(tasks, event, avg_feedback)

    summary_title = document.add_paragraph()
    summary_title.add_run("ИТОГИ").bold = True
    _set_report_font(summary_title, size=13, bold=True)

    employee_title = document.add_paragraph()
    employee_title.add_run("Оценка работы сотрудников").bold = True
    _set_report_font(employee_title)
    employee_table_rows = employee_scores or [("Нет данных", "Н/Д")]
    _add_report_table(document, ["Сотрудник", "Оценка"], employee_table_rows)

    final = document.add_paragraph()
    final.add_run(f"Средняя оценка мероприятия — {avg_feedback}").bold = True
    _set_report_font(final, size=13, bold=True)

    return document


def _build_expense_report_doc(event):
    document = Document()
    orders = _get_event_orders(event)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"ОТЧЕТ О РАСХОДАХ ОРГАНИЗАЦИИ НА МЕРОПРИЯТИЕ № {event.pk}\n").bold = True
    p.add_run((event.title or "").upper()).bold = True
    _set_report_font(p, size=13, bold=True)

    auto = document.add_paragraph()
    auto.add_run("СГЕНЕРИРОВАНО АВТОМАТИЧЕСКИ").italic = True
    _set_report_font(auto, size=12, italic=True)

    info = document.add_paragraph()
    info.add_run("Дата и время проведения мероприятия – ").bold = True
    info.add_run(event.time or "Не указано")
    _set_report_font(info)

    title = document.add_paragraph()
    title.add_run("Затраты").bold = True
    _set_report_font(title)

    rows = []
    total = 0
    for order in orders:
        amount = order.quantity * order.price
        total += amount
        rows.append(
            [
                order.c.name,
                f"ORD-{order.order_id}",
                order.date,
                order.product,
                order.quantity,
                f"{order.price:,.2f}".replace(",", " "),
                f"{amount:,.2f}".replace(",", " "),
            ]
        )
    if rows:
        rows.append(["", "", "", "", "", "ИТОГО", f"{total:,.2f}".replace(",", " ")])
        _add_report_table(
            document,
            [
                "Юр Лицо",
                "Номер договора",
                "Дата заключения",
                "Товар/услуга",
                "Количество",
                "Стоимость 1 ед. (руб.)",
                "Расходы (руб.)",
            ],
            rows,
        )
    else:
        empty = document.add_paragraph("Затраты по мероприятию отсутствуют.")
        _set_report_font(empty)

    return document


def login_view(request):
    if request.user.is_authenticated:
        return redirect("main_page")

    form = LoginUserForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        user = authenticate(
            request,
            login=form.cleaned_data["login"],
            password=form.cleaned_data["password"],
        )
        if user is not None and user.is_active:
            login(request, user)
            return redirect("main_page")
        messages.error(request, "Неверный логин или пароль.")

    return render(request, "main/login.html", {"form": form})


@login_required
def logout_view(request):
    logout(request)
    return redirect("login")


@login_required
def main_page(request):
    events_qs = _visible_events_queryset(request.user)
    visible_tasks = _visible_tasks_queryset(request.user)
    recent_events = list(events_qs[:5])
    employee_map = _fetch_event_employee_map([event.pk for event in recent_events])
    for event in recent_events:
        event.employee_names = employee_map.get(event.pk, [])

    context = {
        "events_count": events_qs.count(),
        "participants_count": Participant.objects.filter(event__in=events_qs).count(),
        "tasks_count": visible_tasks.count(),
        "contractors_count": Contractor.objects.count(),
        "places_count": Place.objects.count(),
        "reports_count": Report.objects.filter(event__in=events_qs).count(),
        "expenses_count": Order.objects.filter(event__in=events_qs).count(),
        "messages_count": Message.objects.filter(receiver=request.user, is_read=False).count(),
        "recent_events": recent_events,
        "recent_tasks": visible_tasks[:5],
    }
    return render(request, "main/mainpage.html", context)


@login_required
def contractors(request):
    query = request.GET.get("search", "").strip()
    object_list = Contractor.objects.all()
    if query:
        object_list = object_list.filter(Q(name__icontains=query) | Q(fullname__icontains=query))

    return render(
        request,
        "main/contractors.html",
        {
            "counterparties": object_list,
            "query": query,
        },
    )


@login_required
def add_contractor(request):
    form = ContractorForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Контрагент добавлен.")
        return redirect("contractors")
    return render(request, "main/contractor_form.html", {"form": form, "title": "Новый контрагент"})


@login_required
def edit_contractor(request, pk):
    contractor = get_object_or_404(Contractor, pk=pk)
    form = ContractorForm(request.POST or None, instance=contractor)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Данные контрагента обновлены.")
        return redirect("contractors")
    return render(request, "main/contractor_form.html", {"form": form, "title": "Редактирование контрагента"})


@login_required
def places(request):
    query = request.GET.get("search", "").strip()
    object_list = Place.objects.select_related("c")
    if query:
        object_list = object_list.filter(Q(address__icontains=query) | Q(description__icontains=query))
    return render(request, "main/places.html", {"places": object_list, "query": query})


@login_required
def add_place(request):
    form = PlaceForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Площадка добавлена.")
        return redirect("places")
    return render(request, "main/place_form.html", {"form": form, "title": "Новая площадка"})


@login_required
def edit_place(request, pk):
    place = get_object_or_404(Place, pk=pk)
    form = PlaceForm(request.POST or None, instance=place)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Данные площадки обновлены.")
        return redirect("places")
    return render(request, "main/place_form.html", {"form": form, "title": "Редактирование площадки"})


@login_required
@require_POST
def delete_place(request, pk):
    place = get_object_or_404(Place, pk=pk)
    try:
        with transaction.atomic():
            Event.objects.filter(p_id=place.pk).update(p=None)
            place.delete()
        messages.success(request, "Площадка удалена.")
    except Exception:
        messages.error(request, "Не удалось удалить площадку.")
    return redirect("places")


@login_required
@require_POST
def delete_contractor(request, pk):
    contractor = get_object_or_404(Contractor, pk=pk)
    try:
        with transaction.atomic():
            Place.objects.filter(c_id=contractor.pk).update(c=None)
            Order.objects.filter(c_id=contractor.pk).delete()
            contractor.delete()
        messages.success(request, "Контрагент удален.")
    except Exception:
        messages.error(request, "Не удалось удалить контрагента.")
    return redirect("contractors")


@login_required
def events(request):
    query = request.GET.get("search", "").strip()
    object_list = _visible_events_queryset(request.user)
    if query:
        object_list = object_list.filter(Q(title__icontains=query) | Q(description__icontains=query))

    events_list = list(object_list)
    employee_map = _fetch_event_employee_map([event.pk for event in events_list])
    participant_counts = {
        item["event"]: item["count"]
        for item in Participant.objects.values("event").order_by().annotate(count=Count("participant_id"))
    }
    task_counts = {
        item["event"]: item["count"]
        for item in Task.objects.values("event").order_by().annotate(count=Count("task_id"))
    }
    for event in events_list:
        event.employee_names = employee_map.get(event.pk, [])
        event.participant_count = participant_counts.get(event.pk, 0)
        event.task_count = task_counts.get(event.pk, 0)
        event.planned_budget_value = _planned_budget(event)
        event.expense_total = _expense_total(_get_event_orders(event))
        event.budget_deviation_percent = _budget_deviation_percent(event.planned_budget_value, event.expense_total)

    return render(request, "main/events.html", {"events": events_list, "query": query})


@login_required
def add_event(request):
    form = EventForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        event = form.save(commit=False)
        event.created_by = request.user
        event.save()
        _sync_event_employees(event.pk, list(form.cleaned_data["employees"].values_list("pk", flat=True)))
        messages.success(request, "Мероприятие создано.")
        return redirect("events")
    return render(request, "main/event_form.html", {"form": form, "title": "Новое мероприятие"})


@login_required
def event_detail(request, pk):
    active_tab = request.GET.get("tab", "overview")
    if active_tab not in {"overview", "feedback"}:
        active_tab = "overview"

    event = get_object_or_404(_visible_events_queryset(request.user), pk=pk)
    employee_map = _fetch_event_employee_map([event.pk])
    selected_employee_ids = _fetch_selected_employee_ids(event.pk)
    orders = _get_event_orders(event)
    for expense in orders:
        expense.total_cost = expense.quantity * expense.price
    expense_total = _expense_total(orders)
    planned_budget = _planned_budget(event)
    budget_deviation_percent = _budget_deviation_percent(planned_budget, expense_total)
    tasks = list(_visible_tasks_queryset(request.user).filter(event=event))
    feedbacks = list(Feedback.objects.filter(event=event))
    feedback_keywords = _extract_feedback_keywords(feedbacks)
    rating_histogram = _build_rating_histogram(feedbacks)
    participants_count = Participant.objects.filter(event=event).count()
    attended_count = Participant.objects.filter(event=event, attended=True).count()
    attendance_ratio = round((attended_count / participants_count) * 100, 2) if participants_count else 0
    binding_form = EventEmployeeBindingForm(initial={"employees": selected_employee_ids})
    context = {
        "event": event,
        "active_tab": active_tab,
        "employees": employee_map.get(event.pk, []),
        "binding_form": binding_form,
        "selected_employee_ids": selected_employee_ids,
        "participants_count": participants_count,
        "attended_count": attended_count,
        "attendance_ratio": attendance_ratio,
        "is_event_finished": event.status == Event.STATUS_FINISHED,
        "tasks_count": len(tasks),
        "feedbacks": feedbacks,
        "feedback_count": len(feedbacks),
        "feedback_keywords": feedback_keywords,
        "rating_histogram": rating_histogram,
        "rating_chart_data": _build_rating_chart_data(rating_histogram),
        "expenses": orders,
        "planned_budget": planned_budget,
        "expense_total": expense_total,
        "budget_deviation_percent": budget_deviation_percent,
    }
    return render(request, "main/event_detail.html", context)


@login_required
@require_POST
def close_event(request, pk):
    if not request.user.can_close_events:
        return HttpResponseForbidden("Недостаточно прав.")
    event = get_object_or_404(_visible_events_queryset(request.user), pk=pk)
    event.status = Event.STATUS_FINISHED
    if not event.end_date:
        event.end_date = _current_local_datetime_text()
    event.save(update_fields=["status", "end_date"])
    messages.success(request, "Мероприятие закрыто.")
    return redirect("event_detail", pk=event.pk)


@login_required
@require_POST
def bind_event_employees(request, pk):
    event = get_object_or_404(_visible_events_queryset(request.user), pk=pk)
    form = EventEmployeeBindingForm(request.POST)
    if form.is_valid():
        employee_ids = list(form.cleaned_data["employees"].values_list("pk", flat=True))
        _sync_event_employees(event.pk, employee_ids)
        messages.success(request, "Пользователи привязаны к мероприятию.")
    else:
        messages.error(request, "Не удалось сохранить привязку пользователей.")
    return redirect("event_detail", pk=event.pk)


@login_required
def edit_event(request, pk):
    event = get_object_or_404(_visible_events_queryset(request.user), pk=pk)
    initial = {"employees": _fetch_selected_employee_ids(event.pk)}
    form = EventForm(request.POST or None, instance=event, initial=initial)
    if request.method == "POST" and form.is_valid():
        event = form.save()
        _sync_event_employees(event.pk, list(form.cleaned_data["employees"].values_list("pk", flat=True)))
        messages.success(request, "Мероприятие обновлено.")
        return redirect("events")
    return render(request, "main/event_form.html", {"form": form, "title": "Редактирование мероприятия"})


@login_required
@require_POST
def delete_event(request, pk):
    event = get_object_or_404(_visible_events_queryset(request.user), pk=pk)
    report_paths = list(
        Report.objects.filter(event_id=pk).values_list("rep_path", flat=True)
    )

    try:
        with transaction.atomic():
            Task.objects.filter(event_id=pk).delete()
            Participant.objects.filter(event_id=pk).delete()
            Feedback.objects.filter(event_id=pk).delete()
            Order.objects.filter(event_id=pk).delete()
            Report.objects.filter(event_id=pk).delete()
            with connection.cursor() as cursor:
                cursor.execute("DELETE FROM employee_on_event WHERE event_id = %s", [pk])
            event.delete()

        messages.success(request, "Мероприятие удалено.")
    except Exception:
        messages.error(
            request,
            "Не удалось удалить мероприятие вместе со связанными данными.",
        )
        return redirect("events")

    for rep_path in report_paths:
        # The event is already deleted from the database; a locked file
        # should not turn the whole operation into a visible failure.
        _delete_report_file(rep_path)

    return redirect("events")


@login_required
def participants(request):
    query = request.GET.get("search", "").strip()
    event_id = request.GET.get("event", "").strip()
    visible_events = _visible_events_queryset(request.user)
    object_list = Participant.objects.select_related("event").filter(event__in=visible_events)

    if query:
        object_list = object_list.filter(Q(fullname__icontains=query) | Q(email__icontains=query) | Q(phone__icontains=query))
    if event_id:
        object_list = object_list.filter(event_id=event_id)

    return render(
        request,
        "main/participants.html",
        {
            "participants": object_list,
            "events": visible_events,
            "query": query,
            "selected_event": event_id,
        },
    )


@login_required
@require_POST
def update_participant_attendance(request, pk):
    participant = get_object_or_404(
        Participant.objects.filter(event__in=_visible_events_queryset(request.user)),
        pk=pk,
    )
    participant.attended = request.POST.get("attended") == "on"
    participant.save(update_fields=["attended"])

    next_url = request.POST.get("next") or reverse("participants")
    if not url_has_allowed_host_and_scheme(next_url, allowed_hosts={request.get_host()}):
        next_url = reverse("participants")
    return redirect(next_url)


@login_required
def expenses(request):
    query = request.GET.get("search", "").strip()
    event_id = request.GET.get("event", "").strip()
    visible_events = _visible_events_queryset(request.user)
    object_list = Order.objects.select_related("event", "c").filter(event__in=visible_events)
    if query:
        object_list = object_list.filter(Q(product__icontains=query) | Q(c__name__icontains=query))
    if event_id:
        object_list = object_list.filter(event_id=event_id)

    expense_rows = list(object_list)
    total = 0
    for expense in expense_rows:
        expense.total_cost = expense.quantity * expense.price
        total += expense.total_cost

    return render(
        request,
        "main/expenses.html",
        {
            "expenses": expense_rows,
            "events": visible_events,
            "query": query,
            "selected_event": event_id,
            "total": total,
        },
    )


@login_required
def add_expense(request):
    initial = {}
    if request.GET.get("event"):
        initial["event"] = request.GET["event"]
    form = OrderForm(request.POST or None, initial=initial, event_queryset=_visible_events_queryset(request.user))
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Расход добавлен.")
        return redirect("expenses")
    return render(request, "main/expense_form.html", {"form": form, "title": "Новый расход"})


@login_required
def edit_expense(request, pk):
    expense = get_object_or_404(Order.objects.filter(event__in=_visible_events_queryset(request.user)), pk=pk)
    form = OrderForm(request.POST or None, instance=expense, event_queryset=_visible_events_queryset(request.user))
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Расход обновлен.")
        return redirect("expenses")
    return render(request, "main/expense_form.html", {"form": form, "title": "Редактирование расхода"})


@login_required
@require_POST
def delete_expense(request, pk):
    expense = get_object_or_404(Order.objects.filter(event__in=_visible_events_queryset(request.user)), pk=pk)
    expense.delete()
    messages.success(request, "Расход удален.")
    return redirect("expenses")


@login_required
def messages_page(request):
    form = MessageForm(request.POST or None, sender=request.user)
    if request.method == "POST" and form.is_valid():
        message = form.save(commit=False)
        message.sender = request.user
        message.is_read = False
        message.sent_at = timezone.now()
        message.save()
        messages.success(request, "Сообщение отправлено.")
        return redirect("messages")

    inbox = Message.objects.filter(receiver=request.user).select_related("sender", "receiver")
    sent = Message.objects.filter(sender=request.user).select_related("sender", "receiver")
    return render(
        request,
        "main/messages.html",
        {
            "form": form,
            "inbox": inbox,
            "sent": sent,
            "unread_count": inbox.filter(is_read=False).count(),
        },
    )


@login_required
def message_detail(request, pk):
    message_obj = get_object_or_404(
        Message.objects.select_related("sender", "receiver"),
        pk=pk,
    )
    if request.user.pk not in {message_obj.sender_id, message_obj.receiver_id}:
        return HttpResponseForbidden("Недостаточно прав.")

    if message_obj.receiver_id == request.user.pk and not message_obj.is_read:
        message_obj.is_read = True
        message_obj.save(update_fields=["is_read"])

    reply_form = MessageForm(
        request.POST or None,
        sender=request.user,
        initial={
            "receiver": message_obj.sender_id if message_obj.sender_id != request.user.pk else message_obj.receiver_id,
            "subject": f"Re: {message_obj.subject or 'Без темы'}",
        },
    )
    if request.method == "POST" and reply_form.is_valid():
        reply = reply_form.save(commit=False)
        reply.sender = request.user
        reply.is_read = False
        reply.sent_at = timezone.now()
        reply.save()
        messages.success(request, "Ответ отправлен.")
        return redirect("message_detail", pk=message_obj.pk)

    return render(
        request,
        "main/message_detail.html",
        {
            "message_obj": message_obj,
            "reply_form": reply_form,
        },
    )


@login_required
@require_POST
def delete_message(request, pk):
    message_obj = get_object_or_404(Message, pk=pk)
    if request.user.pk not in {message_obj.sender_id, message_obj.receiver_id}:
        return HttpResponseForbidden("Недостаточно прав.")

    message_obj.delete()
    messages.success(request, "Сообщение удалено.")
    return redirect("messages")


@login_required
def add_participant(request):
    initial = {}
    if request.GET.get("event"):
        initial["event"] = request.GET["event"]
    form = ParticipantForm(request.POST or None, initial=initial, event_queryset=_visible_events_queryset(request.user))
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Участник добавлен.")
        return redirect("participants")
    return render(request, "main/participant_form.html", {"form": form, "title": "Новый участник"})


@login_required
def upload_participants_csv(request):
    initial = {}
    if request.GET.get("event"):
        initial["event"] = request.GET["event"]
    form = ParticipantCsvUploadForm(
        request.POST or None,
        request.FILES or None,
        initial=initial,
        event_queryset=_visible_events_queryset(request.user),
    )
    if request.method == "POST" and form.is_valid():
        created_count = _import_participants_from_csv(form.cleaned_data["event"], form.cleaned_data["file"])
        messages.success(request, f"Импортировано участников: {created_count}.")
        return redirect("participants")
    return render(
        request,
        "main/participant_csv_upload.html",
        {"form": form, "title": "Импорт участников из CSV"},
    )


@login_required
def edit_participant(request, pk):
    participant = get_object_or_404(
        Participant.objects.filter(event__in=_visible_events_queryset(request.user)),
        pk=pk,
    )
    form = ParticipantForm(
        request.POST or None,
        instance=participant,
        event_queryset=_visible_events_queryset(request.user),
    )
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Данные участника обновлены.")
        return redirect("participants")
    return render(request, "main/participant_form.html", {"form": form, "title": "Редактирование участника"})


@login_required
@require_POST
def delete_participant(request, pk):
    participant = get_object_or_404(
        Participant.objects.filter(event__in=_visible_events_queryset(request.user)),
        pk=pk,
    )
    participant.delete()
    messages.success(request, "Участник удален.")
    return redirect("participants")


@login_required
def upload_feedback_csv(request):
    initial = {}
    if request.GET.get("event"):
        initial["event"] = request.GET["event"]
    form = FeedbackCsvUploadForm(
        request.POST or None,
        request.FILES or None,
        initial=initial,
        event_queryset=_visible_events_queryset(request.user),
    )
    if request.method == "POST" and form.is_valid():
        event = form.cleaned_data["event"]
        created_count = _import_feedback_from_csv(event, form.cleaned_data["file"])
        messages.success(request, f"Импортировано отзывов: {created_count}.")
        return redirect(f"{reverse('event_detail', kwargs={'pk': event.pk})}?tab=feedback")
    return render(
        request,
        "main/feedback_csv_upload.html",
        {"form": form, "title": "Импорт отзывов из CSV"},
    )


@login_required
@require_POST
def delete_feedback(request, pk):
    feedback = get_object_or_404(
        Feedback.objects.filter(event__in=_visible_events_queryset(request.user)),
        pk=pk,
    )
    event_id = feedback.event_id
    feedback.delete()
    messages.success(request, "Отзыв удален.")
    return redirect(f"{reverse('event_detail', kwargs={'pk': event_id})}?tab=feedback")


@login_required
def tasks(request):
    query = request.GET.get("search", "").strip()
    event_id = request.GET.get("event", "").strip()
    visible_events = _visible_events_queryset(request.user)
    object_list = _visible_tasks_queryset(request.user)
    if query:
        object_list = object_list.filter(Q(title__icontains=query) | Q(description__icontains=query))
    if event_id:
        object_list = object_list.filter(event_id=event_id)
    tasks_list = list(object_list)
    for task in tasks_list:
        task.closed_late = _task_closed_late(task)

    return render(
        request,
        "main/tasks.html",
        {
            "tasks": tasks_list,
            "events": visible_events,
            "query": query,
            "selected_event": event_id,
        },
    )


@login_required
def add_task(request):
    initial = {}
    if request.GET.get("event"):
        initial["event"] = request.GET["event"]
    form = TaskForm(
        request.POST or None,
        operator=request.user,
        actor=request.user,
        event_queryset=_visible_events_queryset(request.user),
        initial=initial,
    )
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Задача создана.")
        return redirect("tasks")
    return render(request, "main/task_form.html", {"form": form, "title": "Новая задача"})


@login_required
def edit_task(request, pk):
    task = get_object_or_404(_visible_tasks_queryset(request.user), pk=pk)
    form = TaskForm(
        request.POST or None,
        instance=task,
        operator=task.operator or request.user,
        actor=request.user,
        event_queryset=_visible_events_queryset(request.user),
    )
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Задача обновлена.")
        return redirect("tasks")
    return render(request, "main/task_form.html", {"form": form, "title": "Редактирование задачи"})


@login_required
@require_POST
def close_task(request, pk):
    if not request.user.can_close_tasks:
        return HttpResponseForbidden("Недостаточно прав.")
    task = get_object_or_404(_visible_tasks_queryset(request.user), pk=pk)
    task.status = Task.STATUS_DONE
    if not task.closed_at:
        task.closed_at = timezone.now()
    task.save(update_fields=["status", "closed_at"])
    messages.success(request, "Задача закрыта.")
    return redirect("tasks")


@login_required
@require_POST
def delete_task(request, pk):
    if not request.user.can_delete_tasks:
        return HttpResponseForbidden("Недостаточно прав.")
    task = get_object_or_404(Task, pk=pk)
    task.delete()
    messages.success(request, "Задача удалена.")
    return redirect("tasks")


def _build_event_analytics(events):
    rows = []
    for event in events:
        start_at = _event_start_datetime(event)
        duration_hours = _event_duration_hours(event)
        duration_group_key = _event_duration_group_key(event)
        participants_count = Participant.objects.filter(event=event).count()
        attended_count = Participant.objects.filter(event=event, attended=True).count()
        tasks_count = Task.objects.filter(event=event).count()
        completed_tasks_count = Task.objects.filter(event=event, status=Task.STATUS_DONE).count()
        timeliness = _task_timeliness_stats(Task.objects.filter(event=event))
        feedback_count = Feedback.objects.filter(event=event).count()
        avg_rating = _average_event_rating(event)
        avg_rating_display = avg_rating if avg_rating is not None else "Н/Д"
        attendance_ratio_value = (attended_count / participants_count) if participants_count else 0
        planned_budget = _planned_budget(event)
        expense_total = _expense_total(_get_event_orders(event))
        budget_deviation_percent = _budget_deviation_percent(planned_budget, expense_total)
        budget_quality_deviation = _budget_quality_deviation(planned_budget, expense_total)
        complex_score = _complex_event_score(avg_rating, attendance_ratio_value, planned_budget, expense_total)
        rows.append(
            {
                "event": event,
                "event_date": start_at.strftime("%d.%m.%Y") if start_at else "Н/Д",
                "duration_hours": duration_hours,
                "duration_group_key": duration_group_key,
                "duration_group_label": EVENT_DURATION_GROUP_LABELS[duration_group_key],
                "planned_budget": planned_budget,
                "participants_count": participants_count,
                "attended_count": attended_count,
                "attendance_ratio": attendance_ratio_value,
                "tasks_count": tasks_count,
                "completed_tasks_count": completed_tasks_count,
                "on_time_tasks_percent": timeliness["percent"],
                "feedback_count": feedback_count,
                "avg_rating": avg_rating_display,
                "expense_total": expense_total,
                "budget_deviation_percent": budget_deviation_percent,
                "budget_quality_deviation": round(budget_quality_deviation, 4),
                "quality_score": complex_score,
                "complex_score": complex_score,
            }
        )
    return rows


def _bar_percent(value, max_value):
    if not max_value:
        return 0
    return max(0, min(100, round((value / max_value) * 100)))


def _filter_events_for_analytics(events, period_start=None, period_end=None, duration_group=""):
    filtered = [event for event in events if _event_matches_period(event, period_start, period_end)]
    if duration_group:
        filtered = [event for event in filtered if _event_duration_group_key(event) == duration_group]
    return filtered


def _build_event_quality_chart(event_rows):
    points = []
    for row in event_rows:
        start_at = _event_start_datetime(row["event"])
        if not start_at:
            continue
        score = row["quality_score"]
        points.append(
            {
                "date": start_at.strftime("%d.%m.%Y"),
                "sort": start_at,
                "title": row["event"].title,
                "score": score,
                "bar_percent": max(0, min(100, round(max(score, 0) * 100))),
            }
        )
    return sorted(points, key=lambda item: (item["sort"], item["title"]))


def _build_duration_group_chart(events):
    counts = {key: 0 for key, _label in EVENT_DURATION_GROUPS}
    for event in events:
        counts[_event_duration_group_key(event)] += 1

    max_count = max(counts.values()) if counts else 0
    return [
        {
            "key": key,
            "label": label,
            "count": counts[key],
            "bar_percent": _bar_percent(counts[key], max_count),
        }
        for key, label in EVENT_DURATION_GROUPS
    ]


def _task_month_datetime(task):
    if task.closed_at:
        return _make_aware_local_datetime(task.closed_at)

    deadline_at = _parse_task_deadline(task.deadline)
    if deadline_at:
        return _make_aware_local_datetime(deadline_at)

    return _event_start_datetime(task.event)


def _build_employee_quality_chart(events):
    event_ids = [event.pk for event in events]
    if not event_ids:
        return []

    events_by_id = {event.pk: event for event in events}
    employees_by_id = Employee.objects.in_bulk()
    placeholders = ", ".join(["%s"] * len(event_ids))
    with connection.cursor() as cursor:
        cursor.execute(
            f"SELECT event_id, e_id FROM employee_on_event WHERE event_id IN ({placeholders})",
            event_ids,
        )
        binding_rows = cursor.fetchall()

    grouped = {}
    seen_bindings = set()
    for event_id, employee_id in binding_rows:
        event = events_by_id.get(event_id)
        employee = employees_by_id.get(employee_id)
        if not event or not employee or event.status != Event.STATUS_FINISHED:
            continue
        month_at = _event_start_datetime(event)
        if not month_at:
            continue
        binding_key = (employee_id, event.pk)
        if binding_key in seen_bindings:
            continue
        seen_bindings.add(binding_key)

        month_key = month_at.strftime("%Y-%m")
        key = (employee_id, month_key)
        if key not in grouped:
            grouped[key] = {
                "employee": employee,
                "month": month_at.strftime("%m.%Y"),
                "sort": month_key,
                "event_ids": [],
                "event_quality_scores": [],
            }
        grouped[key]["event_ids"].append(event.pk)
        grouped[key]["event_quality_scores"].append(_event_quality_score(event))

    rows = []
    for row in grouped.values():
        tasks = Task.objects.filter(event_id__in=row["event_ids"], e=row["employee"])
        tasks_count = tasks.count()
        open_tasks_count = tasks.exclude(status=Task.STATUS_DONE).count()
        avg_event_quality = (
            sum(row["event_quality_scores"]) / len(row["event_quality_scores"])
            if row["event_quality_scores"]
            else 0
        )
        unclosed_tasks_ratio = open_tasks_count / tasks_count if tasks_count else 0
        employee_quality_score = avg_event_quality - unclosed_tasks_ratio
        row["events_count"] = len(row["event_ids"])
        row["tasks_count"] = tasks_count
        row["open_tasks_count"] = open_tasks_count
        row["completed_tasks_count"] = tasks_count - open_tasks_count
        row["avg_event_quality"] = round(avg_event_quality, 4)
        row["unclosed_tasks_ratio"] = round(unclosed_tasks_ratio, 4)
        row["employee_quality_score"] = round(employee_quality_score, 4)
        row["productivity_percent"] = round(employee_quality_score * 100, 2)
        row["bar_percent"] = max(0, min(100, round(employee_quality_score * 100)))
        rows.append(row)
    return sorted(rows, key=lambda item: (item["sort"], item["employee"].fullname or item["employee"].login))


def _build_dashboard_summary(event_rows, productivity_rows):
    quality_scores = [row["quality_score"] for row in event_rows]
    employee_quality_scores = [row["employee_quality_score"] for row in productivity_rows]
    participants_count = sum(row["participants_count"] for row in event_rows)
    attended_count = sum(row["attended_count"] for row in event_rows)
    return {
        "events_count": len(event_rows),
        "avg_quality": round(sum(quality_scores) / len(quality_scores), 3) if quality_scores else None,
        "attendance_percent": round((attended_count / participants_count) * 100, 2) if participants_count else None,
        "employee_quality": round(sum(employee_quality_scores) / len(employee_quality_scores), 3) if employee_quality_scores else None,
    }


def _build_analytics_chart_data(quality_chart, duration_groups, productivity_rows):
    month_labels_by_key = {}
    quality_by_employee = {}
    for row in productivity_rows:
        employee_name = row["employee"].fullname or row["employee"].login
        month_labels_by_key[row["sort"]] = row["month"]
        quality_by_employee.setdefault(employee_name, {})[row["sort"]] = row["employee_quality_score"]

    month_keys = sorted(month_labels_by_key)
    return {
        "quality": {
            "labels": [f"{point['date']} · {point['title']}" for point in quality_chart],
            "values": [point["score"] for point in quality_chart],
        },
        "duration": {
            "labels": [group["label"] for group in duration_groups],
            "values": [group["count"] for group in duration_groups],
        },
        "employeeQuality": {
            "labels": [month_labels_by_key[month_key] for month_key in month_keys],
            "datasets": [
                {
                    "label": employee_name,
                    "data": [month_map.get(month_key) for month_key in month_keys],
                }
                for employee_name, month_map in sorted(quality_by_employee.items())
            ],
        },
    }


def _build_employee_analytics(events):
    event_ids = [event.pk for event in events]
    if not event_ids:
        return []

    employees = Employee.objects.filter(employeeonevent__event_id__in=event_ids).distinct()
    rows = []
    for employee in employees:
        related_tasks = Task.objects.filter(event_id__in=event_ids).filter(Q(e=employee) | Q(operator=employee)).distinct()
        rows.append(
            {
                "employee": employee,
                "events_count": Event.objects.filter(event_id__in=event_ids, employeeonevent__e=employee).distinct().count(),
                "tasks_count": related_tasks.count(),
                "completed_tasks_count": related_tasks.filter(status=Task.STATUS_DONE).count(),
            }
        )
    return rows


@login_required
def analytics(request):
    active_tab = request.GET.get("tab", "events")
    if active_tab not in {"events", "employees", "reports"}:
        active_tab = "events"

    date_from = (request.GET.get("date_from") or "").strip()
    date_to = (request.GET.get("date_to") or "").strip()
    period_start = _parse_event_datetime(date_from)
    period_end = _parse_event_datetime(date_to, end_of_day=True)
    selected_duration_group = (request.GET.get("duration_group") or "").strip()
    if selected_duration_group not in EVENT_DURATION_GROUP_LABELS:
        selected_duration_group = ""

    events_list = list(_visible_events_queryset(request.user))
    period_events = _filter_events_for_analytics(events_list, period_start, period_end)
    filtered_events = _filter_events_for_analytics(period_events, duration_group=selected_duration_group)
    event_rows = _build_event_analytics(filtered_events)
    employee_rows = _build_employee_analytics(filtered_events)
    quality_chart = _build_event_quality_chart(event_rows)
    duration_groups = _build_duration_group_chart(period_events)
    employee_quality_chart = _build_employee_quality_chart(filtered_events)
    filter_params = request.GET.copy()
    filter_params.pop("tab", None)
    reports_list = Report.objects.select_related("event").filter(event__in=events_list)
    context = {
        "active_tab": active_tab,
        "event_rows": event_rows,
        "employee_rows": employee_rows,
        "dashboard_summary": _build_dashboard_summary(event_rows, employee_quality_chart),
        "quality_chart": quality_chart,
        "duration_groups": duration_groups,
        "employee_quality_chart": employee_quality_chart,
        "productivity_chart": employee_quality_chart,
        "analytics_chart_data": _build_analytics_chart_data(quality_chart, duration_groups, employee_quality_chart),
        "analytics_filters": {
            "date_from": date_from,
            "date_to": date_to,
            "duration_group": selected_duration_group,
            "duration_group_label": EVENT_DURATION_GROUP_LABELS.get(selected_duration_group, "Все длительности"),
        },
        "filter_query": filter_params.urlencode(),
        "reports": reports_list,
        "events": events_list,
    }
    return render(request, "main/analytics.html", context)


@login_required
def reports(request):
    return redirect("analytics")


def _report_download_response(report_file: Path):
    if not report_file.exists() or not report_file.is_file():
        raise Http404("Файл отчета не найден.")
    return FileResponse(report_file.open("rb"), as_attachment=True, filename=report_file.name)


@login_required
@require_POST
def delete_report(request, pk):
    if not request.user.can_generate_reports:
        return HttpResponseForbidden("Недостаточно прав.")

    report = get_object_or_404(Report, pk=pk)
    rep_path = report.rep_path
    report.delete()
    _delete_report_file(rep_path)
    messages.success(request, "Отчет удален из системы и с компьютера.")
    return redirect(f"{reverse('analytics')}?tab=reports")


@login_required
def download_report(request, pk):
    report = get_object_or_404(
        Report.objects.select_related("event").filter(event__in=_visible_events_queryset(request.user)),
        pk=pk,
    )
    return _report_download_response(Path(report.rep_path))


@login_required
@require_POST
def generate_report(request, event_id, report_type):
    if not request.user.can_generate_reports:
        return HttpResponseForbidden("Недостаточно прав.")

    event = get_object_or_404(_visible_events_queryset(request.user), pk=event_id)
    report_name = Report.TYPE_EVENT if report_type == "event" else Report.TYPE_EXPENSE
    reports_dir = Path(settings.BASE_DIR) / "generated_reports"
    reports_dir.mkdir(exist_ok=True)

    file_name = f"report_{event.pk}_{report_type}_{Report.objects.count() + 1}.docx"
    file_path = reports_dir / file_name
    document = _build_event_report_doc(event) if report_type == "event" else _build_expense_report_doc(event)
    document.save(file_path)

    Report.objects.create(event=event, type=report_name, rep_path=str(file_path))
    return _report_download_response(file_path)


@role_required(Employee.ROLE_ADMIN)
def employees(request):
    query = request.GET.get("search", "").strip()
    object_list = Employee.objects.all()
    if query:
        object_list = object_list.filter(Q(fullname__icontains=query) | Q(login__icontains=query) | Q(email__icontains=query))
    return render(request, "main/employees.html", {"employees": object_list, "query": query})


@role_required(Employee.ROLE_ADMIN)
def add_employee(request):
    form = EmployeeCreationForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Новый пользователь добавлен.")
        return redirect("employees")
    return render(request, "main/employee_form.html", {"form": form, "title": "Новый пользователь"})


@role_required(Employee.ROLE_ADMIN)
def edit_employee(request, pk):
    employee = get_object_or_404(Employee, pk=pk)
    form = EmployeeUpdateForm(request.POST or None, instance=employee)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Данные пользователя обновлены.")
        return redirect("employees")
    return render(
        request,
        "main/employee_form.html",
        {
            "form": form,
            "title": "Редактирование пользователя",
            "is_edit_mode": True,
        },
    )


@role_required(Employee.ROLE_ADMIN)
@require_POST
def delete_employee(request, pk):
    employee = get_object_or_404(Employee, pk=pk)

    if employee.pk == request.user.pk:
        messages.error(request, "Нельзя удалить собственный аккаунт.")
        return redirect("employees")

    try:
        with transaction.atomic():
            Message.objects.filter(Q(sender_id=employee.pk) | Q(receiver_id=employee.pk)).delete()
            Task.objects.filter(e_id=employee.pk).update(e=None)
            Task.objects.filter(operator_id=employee.pk).update(operator=request.user)
            Event.objects.filter(created_by_id=employee.pk).update(created_by=None)
            with connection.cursor() as cursor:
                cursor.execute("DELETE FROM employee_on_event WHERE e_id = %s", [employee.pk])
                cursor.execute("DELETE FROM employee_groups WHERE employee_id = %s", [employee.pk])
                cursor.execute("DELETE FROM employee_user_permissions WHERE employee_id = %s", [employee.pk])
            employee.delete()
        messages.success(request, "Пользователь удален.")
    except Exception:
        messages.error(request, "Не удалось удалить пользователя.")

    return redirect("employees")
