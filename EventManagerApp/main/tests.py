import shutil
import tempfile
import uuid
from io import StringIO
from pathlib import Path

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.db import connection
from django.test import Client, RequestFactory, TestCase, override_settings
from django.utils import timezone
from docx import Document

from .auth_backends import EmployeeAuthBackend
from .decorators import role_required
from .forms import (
    EmployeeCreationForm,
    EmployeeUpdateForm,
    FeedbackCsvUploadForm,
    MessageForm,
    ParticipantCsvUploadForm,
    ParticipantForm,
    TaskForm,
)
from .models import Contractor, Event, Feedback, Message, Order, Participant, Place, Report, Task
from .views import _build_event_analytics, _extract_feedback_keywords


User = get_user_model()


@override_settings(ALLOWED_HOSTS=["testserver", "localhost", "127.0.0.1"])
class EventManagerAppTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        suffix = uuid.uuid4().hex[:8]
        cls.admin_login = f"admin_test_{suffix}"
        cls.manager_login = f"manager_test_{suffix}"
        cls.assistant_login = f"assistant_test_{suffix}"
        cls.teamlead_login = f"teamlead_test_{suffix}"
        cls.legacy_login = f"legacy_test_{suffix}"
        contractor_email = f"contractor_{suffix}@test.local"

        cls.admin = User.objects.create(
            login=cls.admin_login,
            fullname="Администратор Тестовый",
            position=User.ROLE_ADMIN,
            email="admin@test.local",
            is_active=True,
            is_staff=True,
            is_superuser=True,
        )
        cls.admin.set_password("AdminPass123!")
        cls.admin.save(update_fields=["password"])

        cls.manager = User.objects.create(
            login=cls.manager_login,
            fullname="Менеджер Тестовый",
            position=User.ROLE_MANAGER,
            email="manager@test.local",
            is_active=True,
        )
        cls.manager.set_password("ManagerPass123!")
        cls.manager.save(update_fields=["password"])

        cls.assistant = User.objects.create(
            login=cls.assistant_login,
            fullname="Ассистент Тестовый",
            position=User.ROLE_ASSISTANT,
            email="assistant@test.local",
            is_active=True,
        )
        cls.assistant.set_password("AssistantPass123!")
        cls.assistant.save(update_fields=["password"])

        cls.teamlead = User.objects.create(
            login=cls.teamlead_login,
            fullname="Руководитель Тестовый",
            position=User.ROLE_TEAMLEAD,
            email="teamlead@test.local",
            is_active=True,
        )
        cls.teamlead.set_password("TeamleadPass123!")
        cls.teamlead.save(update_fields=["password"])

        cls.legacy_user = User.objects.create(
            login=cls.legacy_login,
            fullname="Legacy Пользователь",
            position=User.ROLE_MANAGER,
            email="legacy@test.local",
            is_active=True,
            password="1234",
        )

        cls.contractor = Contractor.objects.create(
            name='ООО "Тестовый контрагент"',
            fullname="Иванов Иван Иванович",
            email=contractor_email,
            phone="+79990000001",
            type=Contractor.TYPE_RENT_SPACE,
            description="Тестовый контрагент",
        )

    def setUp(self):
        self.client = Client()
        self.temp_dir = tempfile.mkdtemp(prefix="eventmanager-tests-")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def _create_event(self, title="Тестовое мероприятие", created_by=None):
        return Event.objects.create(
            title=title,
            description="Описание тестового мероприятия",
            time="19.04.2026 10:00",
            status=Event.STATUS_DRAFT,
            p=None,
            created_by=created_by,
        )

    def _bind_employee_to_event(self, employee, event):
        with connection.cursor() as cursor:
            cursor.execute(
                "INSERT INTO employee_on_event (e_id, event_id) VALUES (%s, %s)",
                [employee.pk, event.pk],
            )

    def test_employee_auth_backend_handles_edge_cases(self):
        backend = EmployeeAuthBackend()

        self.assertIsNone(backend.authenticate(None, login="", password="pass"))
        self.assertIsNone(backend.authenticate(None, login="missing-user", password="pass"))
        self.assertIsNone(backend.authenticate(None, login=self.manager_login, password="wrong"))

        authenticated = backend.authenticate(None, username=self.manager_login, password="ManagerPass123!")
        self.assertEqual(authenticated, self.manager)
        self.assertEqual(backend.get_user(self.manager.pk), self.manager)
        self.assertIsNone(backend.get_user(999999))

    def test_role_required_rejects_user_without_allowed_role(self):
        request = RequestFactory().get("/protected/")
        request.user = self.assistant

        @role_required(User.ROLE_ADMIN)
        def protected_view(request):
            return "ok"

        with self.assertRaises(PermissionDenied):
            protected_view(request)

    def test_employee_creation_form_validates_login_and_sets_admin_flags(self):
        duplicate_form = EmployeeCreationForm(
            data={
                "fullname": "Duplicate User",
                "login": self.manager_login,
                "email": "duplicate@test.local",
                "phone": "",
                "age": "30",
                "position": User.ROLE_ASSISTANT,
                "is_active": "on",
                "password1": "StrongPass123!",
                "password2": "StrongPass123!",
            }
        )
        self.assertFalse(duplicate_form.is_valid())
        self.assertIn("login", duplicate_form.errors)

        mismatch_form = EmployeeCreationForm(
            data={
                "fullname": "Mismatch User",
                "login": f"mismatch_{uuid.uuid4().hex[:6]}",
                "email": "mismatch@test.local",
                "phone": "",
                "age": "30",
                "position": User.ROLE_ASSISTANT,
                "is_active": "on",
                "password1": "StrongPass123!",
                "password2": "AnotherPass123!",
            }
        )
        self.assertFalse(mismatch_form.is_valid())
        self.assertIn("__all__", mismatch_form.errors)

        admin_login = f"new_admin_{uuid.uuid4().hex[:6]}"
        valid_form = EmployeeCreationForm(
            data={
                "fullname": "New Admin",
                "login": admin_login,
                "email": "new-admin@test.local",
                "phone": "+79990000030",
                "age": "31",
                "position": User.ROLE_ADMIN,
                "is_active": "on",
                "password1": "StrongPass123!",
                "password2": "StrongPass123!",
            }
        )
        self.assertTrue(valid_form.is_valid(), valid_form.errors)
        created = valid_form.save()
        self.assertTrue(created.is_staff)
        self.assertTrue(created.is_superuser)
        self.assertTrue(created.check_password("StrongPass123!"))

    def test_employee_update_form_handles_password_and_duplicate_login(self):
        duplicate_form = EmployeeUpdateForm(
            instance=self.assistant,
            data={
                "fullname": self.assistant.fullname,
                "login": self.manager_login,
                "email": self.assistant.email,
                "phone": self.assistant.phone or "",
                "age": self.assistant.age or "",
                "position": self.assistant.position,
                "is_active": "on",
                "password1": "",
                "password2": "",
            },
        )
        self.assertFalse(duplicate_form.is_valid())
        self.assertIn("login", duplicate_form.errors)

        password_form = EmployeeUpdateForm(
            instance=self.assistant,
            data={
                "fullname": self.assistant.fullname,
                "login": self.assistant.login,
                "email": self.assistant.email,
                "phone": self.assistant.phone or "",
                "age": self.assistant.age or "",
                "position": User.ROLE_TEAMLEAD,
                "is_active": "on",
                "password1": "UpdatedPass123!",
                "password2": "UpdatedPass123!",
            },
        )
        self.assertTrue(password_form.is_valid(), password_form.errors)
        updated = password_form.save()
        self.assertEqual(updated.position, User.ROLE_TEAMLEAD)
        self.assertTrue(updated.check_password("UpdatedPass123!"))

    def test_csv_upload_forms_reject_non_csv_files(self):
        event = self._create_event()
        bad_upload = SimpleUploadedFile("data.txt", b"not,csv", content_type="text/plain")

        participant_form = ParticipantCsvUploadForm(data={"event": event.pk}, files={"file": bad_upload})
        self.assertFalse(participant_form.is_valid())
        self.assertIn("file", participant_form.errors)

        bad_feedback_upload = SimpleUploadedFile("feedback.xlsx", b"not,csv", content_type="application/octet-stream")
        feedback_form = FeedbackCsvUploadForm(data={"event": event.pk}, files={"file": bad_feedback_upload})
        self.assertFalse(feedback_form.is_valid())
        self.assertIn("file", feedback_form.errors)

    def test_participant_form_initializes_instance_gender_and_saves(self):
        event = self._create_event()
        participant = Participant.objects.create(
            event=event,
            fullname="Existing Participant",
            gender=Participant.GENDER_FEMALE,
            phone="+79990000031",
            email="existing-participant@test.local",
        )

        form = ParticipantForm(instance=participant)
        self.assertEqual(form.initial["gender"], "female")

        update_form = ParticipantForm(
            instance=participant,
            data={
                "event": event.pk,
                "fullname": "Updated Participant",
                "gender": "male",
                "phone": "+79990000032",
                "email": "updated-participant@test.local",
            },
        )
        self.assertTrue(update_form.is_valid(), update_form.errors)
        updated = update_form.save()
        self.assertEqual(updated.fullname, "Updated Participant")
        self.assertEqual(updated.gender, Participant.GENDER_MALE)

    def test_task_form_sets_and_clears_closed_at(self):
        event = self._create_event()
        self._bind_employee_to_event(self.assistant, event)

        done_form = TaskForm(
            data={
                "event": event.pk,
                "e": self.assistant.pk,
                "title": "Done task from form",
                "description": "",
                "deadline": "20.04.2026",
                "status": Task.STATUS_DONE,
            },
            operator=self.manager,
            actor=self.manager,
        )
        self.assertTrue(done_form.is_valid(), done_form.errors)
        task = done_form.save()
        self.assertIsNotNone(task.closed_at)

        edit_form = TaskForm(
            instance=task,
            data={
                "event": event.pk,
                "e": self.assistant.pk,
                "title": "Reopened task from form",
                "description": "",
                "deadline": "20.04.2026",
                "status": Task.STATUS_IN_PROGRESS,
            },
            operator=self.manager,
            actor=self.manager,
        )
        self.assertTrue(edit_form.is_valid(), edit_form.errors)
        reopened = edit_form.save()
        self.assertIsNone(reopened.closed_at)

    def test_message_form_excludes_sender_from_receivers(self):
        form = MessageForm(sender=self.manager)

        self.assertNotIn(self.manager.pk, list(form.fields["receiver"].queryset.values_list("pk", flat=True)))
        self.assertIn(self.assistant.pk, list(form.fields["receiver"].queryset.values_list("pk", flat=True)))

    def test_seed_analytics_demo_command_is_idempotent_and_uses_existing_employees(self):
        employee_count = User.objects.count()
        contractor_count = Contractor.objects.count()

        call_command("seed_analytics_demo", stdout=StringIO())
        first_event_ids = set(Event.objects.filter(title__startswith="[Демо аналитика]").values_list("pk", flat=True))
        first_task_count = Task.objects.filter(event_id__in=first_event_ids).count()

        call_command("seed_analytics_demo", stdout=StringIO())
        second_event_ids = set(Event.objects.filter(title__startswith="[Демо аналитика]").values_list("pk", flat=True))

        self.assertEqual(User.objects.count(), employee_count)
        self.assertEqual(Contractor.objects.count(), contractor_count)
        self.assertEqual(Event.objects.filter(title__startswith="[Демо аналитика]").count(), 5)
        self.assertEqual(Task.objects.filter(event_id__in=second_event_ids).count(), first_task_count)
        self.assertEqual(Participant.objects.filter(event_id__in=second_event_ids).count(), 330)
        self.assertEqual(Feedback.objects.filter(event_id__in=second_event_ids).count(), 26)
        months = sorted({Event.objects.get(pk=event_id).time[3:10] for event_id in second_event_ids})
        self.assertGreaterEqual(len(months), 5)

    def test_legacy_plaintext_password_is_accepted_and_upgraded(self):
        response = self.client.post(
            "/login/",
            {"login": self.legacy_login, "password": "1234"},
            follow=False,
        )

        self.assertEqual(response.status_code, 302)
        self.legacy_user.refresh_from_db()
        self.assertNotEqual(self.legacy_user.password, "1234")
        self.assertTrue(self.legacy_user.password.startswith("pbkdf2_"))

    def test_participant_can_be_added(self):
        event = self._create_event(created_by=self.manager)
        self.client.force_login(self.admin)

        response = self.client.post(
            "/participants/add/",
            {
                "event": event.pk,
                "fullname": "Новый Участник",
                "gender": "male",
                "phone": "+79990000002",
                "email": "participant@test.local",
            },
            follow=False,
        )

        self.assertEqual(response.status_code, 302)
        participant = Participant.objects.get(event=event, email="participant@test.local")
        self.assertEqual(participant.fullname, "Новый Участник")
        self.assertEqual(participant.gender, Participant.GENDER_MALE)
        self.assertFalse(participant.attended)

    def test_participant_attendance_can_be_toggled_from_list(self):
        event = self._create_event()
        participant = Participant.objects.create(
            event=event,
            fullname="Участник для отметки",
            email="attendance@test.local",
            attended=False,
        )
        self.client.force_login(self.admin)

        response = self.client.post(
            f"/participants/{participant.pk}/attendance/",
            {"attended": "on", "next": "/participants/"},
            follow=False,
        )

        self.assertEqual(response.status_code, 302)
        participant.refresh_from_db()
        self.assertTrue(participant.attended)

    def test_participants_can_be_imported_from_csv(self):
        event = self._create_event()
        self.client.force_login(self.admin)
        upload = SimpleUploadedFile(
            "participants.csv",
            "fullname,gender,phone,email,attended\nCSV One,male,+79990000011,one@test.local,1\nCSV Two,female,+79990000012,two@test.local,0\n".encode(),
            content_type="text/csv",
        )

        response = self.client.post(
            "/participants/upload-csv/",
            {"event": event.pk, "file": upload},
            follow=False,
        )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(Participant.objects.filter(event=event, email__endswith="@test.local").count(), 2)
        self.assertEqual(Participant.objects.get(event=event, email="one@test.local").gender, Participant.GENDER_MALE)
        self.assertEqual(Participant.objects.get(event=event, email="one@test.local").phone, "+79990000011")
        self.assertFalse(Participant.objects.get(event=event, email="one@test.local").attended)

    def test_participants_can_be_imported_with_phone_number_header(self):
        event = self._create_event(created_by=self.manager)
        self.client.force_login(self.admin)
        upload = SimpleUploadedFile(
            "participants.csv",
            "ФИО,номер телефона,email\nCSV Header,+79990000015,header@test.local\n".encode(),
            content_type="text/csv",
        )

        response = self.client.post(
            "/participants/upload-csv/",
            {"event": event.pk, "file": upload},
            follow=False,
        )

        self.assertEqual(response.status_code, 302)
        participant = Participant.objects.get(event=event, email="header@test.local")
        self.assertEqual(participant.phone, "+79990000015")

    def test_event_can_be_closed_and_attendance_ratio_is_displayed(self):
        event = self._create_event(created_by=self.manager)
        Participant.objects.create(event=event, fullname="Пришел", attended=True)
        Participant.objects.create(event=event, fullname="Не пришел", attended=False)
        self.client.force_login(self.manager)

        close_response = self.client.post(f"/events/{event.pk}/close/")

        self.assertEqual(close_response.status_code, 302)
        event.refresh_from_db()
        self.assertEqual(event.status, Event.STATUS_FINISHED)
        self.assertTrue(event.end_date)

        detail_response = self.client.get(f"/events/{event.pk}/")
        self.assertContains(detail_response, "1 / 2")

    def test_feedback_can_be_imported_from_csv(self):
        event = self._create_event()
        self.client.force_login(self.admin)
        upload = SimpleUploadedFile(
            "feedback.csv",
            "оценка,отзыв\n9,Очень хорошо\n7,Нужно больше навигации\n".encode(),
            content_type="text/csv",
        )

        response = self.client.post(
            "/feedback/upload-csv/",
            {"event": event.pk, "file": upload},
            follow=False,
        )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(Feedback.objects.filter(event=event).count(), 2)
        self.assertEqual(Feedback.objects.get(event=event, rating=9).review, "Очень хорошо")

    def test_feedback_import_supports_survey_export_headers(self):
        event = self._create_event()
        self.client.force_login(self.admin)
        upload = SimpleUploadedFile(
            "feedback.csv",
            (
                "Введите ваше ФИО;Оцените прошедшее мероприятие от 1 до 10;Расскажите что вам понравилось/не понравилось на прошедшем мероприятии\n"
                "Дорофеев Дмитрий Сергеевич;6;Очень долгая регистрация\n"
                "Участник без имени;2;Неинтересно\n"
            ).encode(),
            content_type="text/csv",
        )

        response = self.client.post(
            "/feedback/upload-csv/",
            {"event": event.pk, "file": upload},
            follow=False,
        )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(Feedback.objects.filter(event=event).count(), 2)
        self.assertEqual(Feedback.objects.get(event=event, rating=6).review, "Очень долгая регистрация")

    def test_event_detail_shows_yake_keywords_from_feedback(self):
        event = self._create_event()
        Feedback.objects.create(
            event=event,
            rating=6,
            review="долгая регистрация долгая регистрация перед мероприятием",
        )
        Feedback.objects.create(
            event=event,
            rating=8,
            review="понравилась организация, но регистрация заняла много времени",
        )
        self.client.force_login(self.admin)

        response = self.client.get(f"/events/{event.pk}/?tab=feedback")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Ключевые слова из отзывов")
        self.assertContains(response, "долгая регистрация")

    def test_feedback_keywords_do_not_repeat_nested_terms(self):
        event = self._create_event()
        feedbacks = [
            Feedback(event=event, rating=6, review="долгая регистрация долгая регистрация"),
            Feedback(event=event, rating=8, review="регистрация заняла много времени"),
        ]

        keywords = [_keyword.casefold() for _keyword in _extract_feedback_keywords(feedbacks)]

        self.assertIn("долгая регистрация", keywords)
        self.assertNotIn("регистрация", keywords)
        self.assertEqual(len(keywords), len(set(keywords)))

    def test_participant_csv_import_detects_phone_column_without_headers(self):
        event = self._create_event()
        self.client.force_login(self.admin)
        upload = SimpleUploadedFile(
            "participants.csv",
            "CSV Three;89990000013;three@test.local\nCSV Four;79990000014.0;four@test.local\n".encode(),
            content_type="text/csv",
        )

        response = self.client.post(
            "/participants/upload-csv/",
            {"event": event.pk, "file": upload},
            follow=False,
        )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(Participant.objects.get(event=event, email="three@test.local").phone, "+79990000013")
        self.assertEqual(Participant.objects.get(event=event, email="four@test.local").phone, "+79990000014")

    def test_feedback_manual_form_is_removed(self):
        event = self._create_event()
        self.client.force_login(self.admin)

        response = self.client.post(
            "/feedback/add/",
            {
                "event": event.pk,
                "rating": "9",
                "review": "Everything was well organized.",
            },
            follow=False,
        )

        self.assertEqual(response.status_code, 404)
        self.assertFalse(Feedback.objects.filter(event=event).exists())

    def test_event_employee_binding_saves_selected_users(self):
        event = self._create_event()
        self.client.force_login(self.admin)

        response = self.client.post(
            f"/events/{event.pk}/bindings/",
            {"employees": [self.manager.pk, self.assistant.pk]},
            follow=False,
        )

        self.assertEqual(response.status_code, 302)
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT e_id FROM employee_on_event WHERE event_id = %s ORDER BY e_id",
                [event.pk],
            )
            rows = [row[0] for row in cursor.fetchall()]
        self.assertEqual(rows, sorted([self.manager.pk, self.assistant.pk]))

    def test_event_list_is_limited_by_role_creator_and_responsibility(self):
        manager_event = self._create_event("manager-created-event", created_by=self.manager)
        assigned_event = self._create_event("assigned-event")
        hidden_event = self._create_event("hidden-event")
        self._bind_employee_to_event(self.manager, assigned_event)
        self._bind_employee_to_event(self.assistant, assigned_event)

        self.client.force_login(self.manager)
        manager_response = self.client.get("/events/")
        self.assertContains(manager_response, manager_event.title)
        self.assertContains(manager_response, assigned_event.title)
        self.assertNotContains(manager_response, hidden_event.title)

        self.client.force_login(self.assistant)
        assistant_response = self.client.get("/events/")
        self.assertContains(assistant_response, assigned_event.title)
        self.assertNotContains(assistant_response, manager_event.title)
        self.assertNotContains(assistant_response, hidden_event.title)

        self.client.force_login(self.teamlead)
        teamlead_response = self.client.get("/events/")
        self.assertContains(teamlead_response, manager_event.title)
        self.assertContains(teamlead_response, assigned_event.title)
        self.assertContains(teamlead_response, hidden_event.title)

    def test_event_analytics_calculates_budget_deviation_and_quality_score(self):
        event = self._create_event("analytics-budget-event")
        event.planned_budget = 12000
        event.save(update_fields=["planned_budget"])
        Participant.objects.create(event=event, fullname="Пришел", attended=True)
        Participant.objects.create(event=event, fullname="Не пришел", attended=False)
        Feedback.objects.create(event=event, rating=8, review="Хорошо")
        Order.objects.create(
            event=event,
            c=self.contractor,
            product="Аренда",
            quantity=1,
            price=10000,
            date="19.04.2026",
        )

        row = _build_event_analytics([event])[0]

        self.assertEqual(row["planned_budget"], 12000)
        self.assertEqual(row["expense_total"], 10000)
        self.assertEqual(row["budget_deviation_percent"], -16.67)
        self.assertEqual(row["budget_quality_deviation"], 0.1667)
        self.assertEqual(row["quality_score"], 0.7633)

    def test_analytics_filters_events_by_period_and_duration_group(self):
        short_event = self._create_event("short-duration-dashboard-event")
        short_event.time = "2026-04-20 10:00"
        short_event.end_date = "2026-04-20 12:00"
        short_event.save(update_fields=["time", "end_date"])

        long_event = self._create_event("long-duration-dashboard-event")
        long_event.time = "2026-04-21 10:00"
        long_event.end_date = "2026-04-24 10:00"
        long_event.save(update_fields=["time", "end_date"])

        outside_event = self._create_event("outside-dashboard-event")
        outside_event.time = "2026-05-20 10:00"
        outside_event.end_date = "2026-05-20 12:00"
        outside_event.save(update_fields=["time", "end_date"])

        self.client.force_login(self.admin)
        response = self.client.get(
            "/analytics/",
            {
                "tab": "events",
                "date_from": "2026-04-19",
                "date_to": "2026-04-22",
                "duration_group": "short",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Качество мероприятий по датам")
        self.assertContains(response, "eventQualityChart")
        self.assertContains(response, "durationGroupChart")
        self.assertContains(response, "chart.umd.min.js")
        self.assertContains(response, "Короткие до 3 часов")
        self.assertContains(response, short_event.title)
        self.assertNotContains(response, long_event.title)
        self.assertNotContains(response, outside_event.title)

    def test_employee_productivity_dashboard_groups_tasks_by_month(self):
        event = self._create_event("employee-productivity-dashboard-event")
        event.time = "2026-03-01 10:00"
        event.end_date = "2026-03-01 14:00"
        event.status = Event.STATUS_FINISHED
        event.save(update_fields=["time", "end_date", "status"])
        self._bind_employee_to_event(self.assistant, event)
        Participant.objects.create(event=event, fullname="Пришел", attended=True)
        Feedback.objects.create(event=event, rating=8, review="Хорошо")
        Task.objects.create(
            event=event,
            operator=self.manager,
            e=self.assistant,
            title="completed-march-task",
            description="",
            deadline="10.03.2026",
            status=Task.STATUS_DONE,
            closed_at=timezone.datetime(2026, 3, 10, 12, 0, tzinfo=timezone.get_current_timezone()),
        )
        Task.objects.create(
            event=event,
            operator=self.manager,
            e=self.assistant,
            title="open-march-task",
            description="",
            deadline="20.03.2026",
            status=Task.STATUS_NEW,
        )

        self.client.force_login(self.admin)
        response = self.client.get("/analytics/", {"tab": "employees"})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Качество работы сотрудников по месяцам")
        self.assertContains(response, "employeeQualityChart")
        self.assertContains(response, self.assistant.fullname)
        self.assertContains(response, "03.2026")

    def test_event_detail_shows_rating_histogram(self):
        event = self._create_event("rating-histogram-event")
        Feedback.objects.create(event=event, rating=6, review="Нормально")
        Feedback.objects.create(event=event, rating=6, review="Можно лучше")
        Feedback.objects.create(event=event, rating=9, review="Отлично")
        self.client.force_login(self.admin)

        response = self.client.get(f"/events/{event.pk}/", {"tab": "feedback"})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Распределение оценок")
        self.assertContains(response, "eventRatingChart")
        self.assertContains(response, "6/10")
        self.assertContains(response, "9/10")

    def test_manager_can_assign_task_only_to_linked_assistant(self):
        event = self._create_event("manager-task-event", created_by=self.manager)
        self._bind_employee_to_event(self.assistant, event)
        self._bind_employee_to_event(self.teamlead, event)
        self.client.force_login(self.manager)

        success_response = self.client.post(
            "/tasks/add/",
            {
                "event": event.pk,
                "e": self.assistant.pk,
                "title": "assistant-task",
                "description": "",
                "deadline": "20.04.2026",
                "status": Task.STATUS_NEW,
            },
            follow=False,
        )

        self.assertEqual(success_response.status_code, 302)
        self.assertTrue(Task.objects.filter(event=event, e=self.assistant, title="assistant-task").exists())

        forbidden_response = self.client.post(
            "/tasks/add/",
            {
                "event": event.pk,
                "e": self.teamlead.pk,
                "title": "teamlead-task",
                "description": "",
                "deadline": "20.04.2026",
                "status": Task.STATUS_NEW,
            },
            follow=False,
        )

        self.assertEqual(forbidden_response.status_code, 200)
        self.assertFalse(Task.objects.filter(event=event, title="teamlead-task").exists())

    def test_task_assignee_must_be_linked_to_event(self):
        event = self._create_event("linked-assignee-event")
        self.client.force_login(self.admin)

        response = self.client.post(
            "/tasks/add/",
            {
                "event": event.pk,
                "e": self.assistant.pk,
                "title": "unlinked-assistant-task",
                "description": "",
                "deadline": "20.04.2026",
                "status": Task.STATUS_NEW,
            },
            follow=False,
        )

        self.assertEqual(response.status_code, 200)
        self.assertFalse(Task.objects.filter(event=event, title="unlinked-assistant-task").exists())

    def test_message_send_and_read_flow(self):
        self.client.force_login(self.manager)
        send_response = self.client.post(
            "/messages/",
            {
                "receiver": self.assistant.pk,
                "subject": "Тестовая тема",
                "body": "Тестовое сообщение",
            },
            follow=False,
        )

        self.assertEqual(send_response.status_code, 302)
        message = Message.objects.get(subject="Тестовая тема")
        self.assertFalse(message.is_read)

        self.client.force_login(self.assistant)
        detail_response = self.client.get(f"/messages/{message.pk}/")
        self.assertEqual(detail_response.status_code, 200)
        message.refresh_from_db()
        self.assertTrue(message.is_read)

    def test_message_can_be_deleted_by_participant_of_dialog(self):
        message = Message.objects.create(
            sender=self.manager,
            receiver=self.assistant,
            subject="Удаляемое сообщение",
            body="Сообщение для удаления",
            sent_at=timezone.now(),
            is_read=False,
        )

        self.client.force_login(self.assistant)
        response = self.client.post(f"/messages/{message.pk}/delete/")

        self.assertEqual(response.status_code, 302)
        self.assertFalse(Message.objects.filter(pk=message.pk).exists())

    def test_manager_cannot_delete_task_but_teamlead_can(self):
        event = self._create_event()
        task = Task.objects.create(
            event=event,
            operator=self.manager,
            e=self.manager,
            title="Задача для удаления",
            description="",
            deadline="",
            status=Task.STATUS_NEW,
        )

        self.client.force_login(self.manager)
        forbidden_response = self.client.post(f"/tasks/{task.pk}/delete/")
        self.assertEqual(forbidden_response.status_code, 403)
        self.assertTrue(Task.objects.filter(pk=task.pk).exists())

        self.client.force_login(self.teamlead)
        success_response = self.client.post(f"/tasks/{task.pk}/delete/")
        self.assertEqual(success_response.status_code, 302)
        self.assertFalse(Task.objects.filter(pk=task.pk).exists())

    def test_task_list_is_limited_to_visible_tasks(self):
        event = self._create_event()
        self._bind_employee_to_event(self.assistant, event)
        own_task = Task.objects.create(
            event=event,
            operator=self.manager,
            e=self.assistant,
            title="visible-assistant-task",
            description="",
            deadline="",
            status=Task.STATUS_NEW,
        )
        other_task = Task.objects.create(
            event=event,
            operator=self.admin,
            e=self.manager,
            title="hidden-manager-task",
            description="",
            deadline="",
            status=Task.STATUS_NEW,
        )

        self.client.force_login(self.assistant)
        response = self.client.get("/tasks/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, own_task.title)
        self.assertNotContains(response, other_task.title)

    def test_manager_can_close_created_task_and_closed_at_is_saved(self):
        event = self._create_event(created_by=self.manager)
        task = Task.objects.create(
            event=event,
            operator=self.manager,
            e=self.assistant,
            title="closable-task",
            description="",
            deadline="",
            status=Task.STATUS_NEW,
        )

        self.client.force_login(self.manager)
        response = self.client.post(f"/tasks/{task.pk}/close/")

        self.assertEqual(response.status_code, 302)
        task.refresh_from_db()
        self.assertEqual(task.status, Task.STATUS_DONE)
        self.assertIsNotNone(task.closed_at)

    def test_late_closed_task_is_marked_in_task_list(self):
        event = self._create_event()
        Task.objects.create(
            event=event,
            operator=self.admin,
            e=self.admin,
            title="late-closed-task",
            description="",
            deadline="20.04.2026",
            status=Task.STATUS_DONE,
            closed_at=timezone.datetime(2026, 4, 21, 9, 0, tzinfo=timezone.get_current_timezone()),
        )
        self.client.force_login(self.admin)

        response = self.client.get("/tasks/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "late-closed-task")
        self.assertContains(response, "row-danger")

    def test_delete_report_removes_database_record_and_file(self):
        event = self._create_event()
        report_file = Path(self.temp_dir) / "report.docx"
        report_file.write_text("temporary", encoding="utf-8")
        report = Report.objects.create(
            event=event,
            type=Report.TYPE_EVENT,
            rep_path=str(report_file),
        )
        self.client.force_login(self.admin)

        response = self.client.post(f"/reports/{report.pk}/delete/", follow=False)

        self.assertEqual(response.status_code, 302)
        self.assertFalse(Report.objects.filter(pk=report.pk).exists())
        self.assertFalse(report_file.exists())

    def test_delete_place_clears_event_reference_and_removes_place(self):
        place = Place.objects.create(
            c=self.contractor,
            address="г. Москва, ул. Тестовая, 1",
            description="Тестовая площадка",
        )
        event = Event.objects.create(
            title="Мероприятие с площадкой",
            description="",
            time="19.04.2026 12:00",
            status=Event.STATUS_DRAFT,
            p=place,
        )
        self.client.force_login(self.admin)

        response = self.client.post(f"/places/{place.pk}/delete/")

        self.assertEqual(response.status_code, 302)
        self.assertFalse(Place.objects.filter(pk=place.pk).exists())
        event.refresh_from_db()
        self.assertIsNone(event.p)

    def test_delete_contractor_clears_places_and_removes_related_expenses(self):
        event = self._create_event("Удаление контрагента")
        place = Place.objects.create(
            c=self.contractor,
            address="г. Москва, ул. Контрагентская, 2",
            description="Площадка контрагента",
        )
        expense = Order.objects.create(
            event=event,
            c=self.contractor,
            product="Услуга контрагента",
            quantity=1,
            price=2000.0,
            date="19.04.2026",
        )
        self.client.force_login(self.admin)

        response = self.client.post(f"/contractors/{self.contractor.pk}/delete/")

        self.assertEqual(response.status_code, 302)
        self.assertFalse(Contractor.objects.filter(pk=self.contractor.pk).exists())
        self.assertFalse(Order.objects.filter(pk=expense.pk).exists())
        place.refresh_from_db()
        self.assertIsNone(place.c)

    def test_delete_event_cascades_to_related_entities_and_files(self):
        event = self._create_event("Каскадное удаление")
        Participant.objects.create(
            event=event,
            fullname="Каскадный участник",
            phone="+79990000003",
            email="cascade@test.local",
        )
        Feedback.objects.create(
            event=event,
            rating=8,
            review="Каскадный отзыв",
        )
        Task.objects.create(
            event=event,
            operator=self.manager,
            e=self.manager,
            title="Каскадная задача",
            description="",
            deadline="",
            status=Task.STATUS_NEW,
        )
        Order.objects.create(
            event=event,
            c=self.contractor,
            product="Каскадный расход",
            quantity=2,
            price=1500.0,
            date="19.04.2026",
        )
        report_file = Path(self.temp_dir) / "cascade.docx"
        report_file.write_text("cascade", encoding="utf-8")
        Report.objects.create(
            event=event,
            type=Report.TYPE_EVENT,
            rep_path=str(report_file),
        )
        self._bind_employee_to_event(self.manager, event)
        self.client.force_login(self.admin)

        response = self.client.post(f"/events/{event.pk}/delete/", follow=False)

        self.assertEqual(response.status_code, 302)
        self.assertFalse(Event.objects.filter(pk=event.pk).exists())
        self.assertFalse(Participant.objects.filter(event=event).exists())
        self.assertFalse(Feedback.objects.filter(event=event).exists())
        self.assertFalse(Task.objects.filter(event=event).exists())
        self.assertFalse(Order.objects.filter(event=event).exists())
        self.assertFalse(Report.objects.filter(event=event).exists())
        self.assertFalse(report_file.exists())
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT COUNT(*) FROM employee_on_event WHERE event_id = %s",
                [event.pk],
            )
            self.assertEqual(cursor.fetchone()[0], 0)

    def test_admin_can_delete_other_user_but_not_self(self):
        self.client.force_login(self.admin)

        self.assertTrue(User.objects.filter(pk=self.manager.pk).exists())
        delete_other_response = self.client.post(f"/employees/{self.manager.pk}/delete/")
        self.assertEqual(delete_other_response.status_code, 302)
        self.assertFalse(User.objects.filter(pk=self.manager.pk).exists())

        self.assertTrue(User.objects.filter(pk=self.admin.pk).exists())
        delete_self_response = self.client.post(f"/employees/{self.admin.pk}/delete/")
        self.assertEqual(delete_self_response.status_code, 302)
        self.assertTrue(User.objects.filter(pk=self.admin.pk).exists())

    def test_admin_can_edit_user_data(self):
        self.client.force_login(self.admin)

        response = self.client.post(
            f"/employees/{self.assistant.pk}/edit/",
            {
                "fullname": "Ассистент Обновленный",
                "login": self.assistant.login,
                "email": "assistant-updated@test.local",
                "phone": "+79990000099",
                "age": "31",
                "position": User.ROLE_ASSISTANT,
                "is_active": "on",
                "password1": "",
                "password2": "",
            },
            follow=False,
        )

        self.assertEqual(response.status_code, 302)
        self.assistant.refresh_from_db()
        self.assertEqual(self.assistant.fullname, "Ассистент Обновленный")
        self.assertEqual(self.assistant.email, "assistant-updated@test.local")
        self.assertEqual(self.assistant.phone, "+79990000099")

    def test_report_generation_creates_docx_and_db_record(self):
        event = self._create_event("Отчетное мероприятие")
        Task.objects.create(
            event=event,
            operator=self.teamlead,
            e=self.teamlead,
            title="Подготовить отчет",
            description="",
            deadline="20.04.2026",
            status=Task.STATUS_DONE,
            closed_at=timezone.datetime(2026, 4, 20, 12, 0, tzinfo=timezone.get_current_timezone()),
        )
        Task.objects.create(
            event=event,
            operator=self.teamlead,
            e=self.teamlead,
            title="Незавершенная задача",
            description="",
            deadline="21.04.2026",
            status=Task.STATUS_NEW,
        )
        Feedback.objects.create(
            event=event,
            rating=10,
            review="Отзыв попадает в отчет",
        )
        self._bind_employee_to_event(self.teamlead, event)
        self.client.force_login(self.teamlead)

        original_base_dir = settings.BASE_DIR
        settings.BASE_DIR = Path(self.temp_dir)
        try:
            response = self.client.post(
                f"/reports/generate/{event.pk}/event/",
                follow=False,
            )
            self.assertEqual(response.status_code, 200)
            self.assertIn("attachment;", response["Content-Disposition"])
            report = Report.objects.filter(event=event, type=Report.TYPE_EVENT).latest("rep_id")
            self.assertTrue(Path(report.rep_path).exists())
            self.assertEqual(Path(report.rep_path).suffix.lower(), ".docx")
            document = Document(report.rep_path)
            table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            paragraph_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("Отзыв попадает в отчет", table_text)
            self.assertIn("Процент вовремя закрытых задач", paragraph_text)
            self.assertIn("100.0%", paragraph_text)
            self.assertIn("9.5", table_text)
            self.assertNotIn("Оценка контрагентов", paragraph_text)
        finally:
            settings.BASE_DIR = original_base_dir

    def test_report_generation_downloads_created_docx(self):
        event = self._create_event("Скачиваемый отчет")
        self._bind_employee_to_event(self.teamlead, event)
        self.client.force_login(self.teamlead)

        original_base_dir = settings.BASE_DIR
        settings.BASE_DIR = Path(self.temp_dir)
        try:
            response = self.client.post(f"/reports/generate/{event.pk}/event/", follow=False)
            self.assertEqual(response.status_code, 200)
            self.assertIn("attachment;", response["Content-Disposition"])
            self.assertIn(".docx", response["Content-Disposition"])

            report = Report.objects.filter(event=event, type=Report.TYPE_EVENT).latest("rep_id")
            self.assertTrue(Path(report.rep_path).exists())
            self.assertGreater(len(b"".join(response.streaming_content)), 0)
        finally:
            settings.BASE_DIR = original_base_dir

    def test_download_report_returns_existing_file(self):
        event = self._create_event("Готовый отчет")
        report_file = Path(self.temp_dir) / "existing_report.docx"
        report_file.write_bytes(b"existing report bytes")
        report = Report.objects.create(
            event=event,
            type=Report.TYPE_EVENT,
            rep_path=str(report_file),
        )
        self._bind_employee_to_event(self.teamlead, event)
        self.client.force_login(self.teamlead)

        response = self.client.get(f"/reports/{report.pk}/download/", follow=False)

        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment;", response["Content-Disposition"])
        self.assertIn(report_file.name, response["Content-Disposition"])
        self.assertEqual(b"".join(response.streaming_content), b"existing report bytes")
